from pathlib import Path
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.page import PageMargins
from pptx import Presentation
from pptx.util import Inches as PInches
from PIL import Image, ImageDraw, ImageFont, ImageFilter

ROOT = Path(__file__).resolve().parents[1] / "test" / "fixtures" / "processing"
ROOT.mkdir(parents=True, exist_ok=True)
font_candidates = [
    Path("/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf"),
    *Path("/usr/share/fonts").glob("**/NotoSansArabic-Regular.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
]
font_path = str(next(path for path in font_candidates if path.exists()))
font = ImageFont.truetype(font_path, 44)

image = Image.new("RGB", (800, 300), "white")
draw = ImageDraw.Draw(image); draw.rectangle((20, 20, 780, 280), outline="navy", width=5); draw.text((60, 100), "DIP generated visual", fill="navy", font=ImageFont.truetype(font_path, 34))
visual = ROOT / "visual.png"; image.save(visual)

for name, paragraphs in {
    "word-en.docx": ["DIP generated English document", "Safe representative fixture."],
    "word-ar.docx": ["مستند عربي تجريبي لمنصة معالجة المستندات", "هذا ملف آمن مولد للاختبار"],
    "word-mixed.docx": ["DIP منصة المستندات — mixed Arabic and English"],
}.items():
    doc = Document();
    for text in paragraphs: doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = "Item"; table.cell(0, 1).text = "القيمة"; table.cell(1, 0).text = "DIP"; table.cell(1, 1).text = "١٨"
    doc.add_picture(str(visual), width=Inches(3)); doc.save(ROOT / name)

wb = Workbook(); ws = wb.active; ws.title = "English"; ws.append(["Month", "Jobs"]); ws.append(["Jan", 10]); ws.append(["Feb", 18]); ws.print_area = "A1:B3"; ws.page_margins = PageMargins(left=.25, right=.25, top=.5, bottom=.5)
ar = wb.create_sheet("العربية"); ar.sheet_view.rightToLeft = True; ar.append(["الشهر", "المهام"]); ar.append(["يناير", 12])
chart = BarChart(); chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True); chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3)); ws.add_chart(chart, "D2")
wb.save(ROOT / "excel-corpus.xlsx")

prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text = "DIP English"; slide.placeholders[1].text = "Generated common layout"
slide = prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text = "منصة ذكاء المستندات"; slide.shapes.add_picture(str(visual), PInches(1), PInches(2), width=PInches(5)); prs.save(ROOT / "powerpoint-corpus.pptx")

def scan(name, lines, blur=False, pages=1):
    imgs=[]
    for p in range(pages):
        img=Image.new("RGB", (1240, 1754), "white"); d=ImageDraw.Draw(img); y=180
        for line in lines: d.text((100, y), line, fill="black", font=font); y += 100
        if blur: img=img.filter(ImageFilter.GaussianBlur(1.2))
        imgs.append(img)
    path=ROOT / name; imgs[0].save(path, save_all=True, append_images=imgs[1:], resolution=150)

scan("ocr-en.pdf", ["Generated English scan", "Document intelligence platform"])
scan("ocr-ar.pdf", ["مستند عربي مولد للاختبار", "منصة ذكاء المستندات"])
scan("ocr-mixed.pdf", ["DIP Document", "منصة المستندات العربية English"])
scan("ocr-low-quality.pdf", ["Low quality generated scan", "اختبار جودة منخفضة"], blur=True)
scan("ocr-multipage.pdf", ["Multi page OCR fixture", "اختبار متعدد الصفحات"], pages=2)
